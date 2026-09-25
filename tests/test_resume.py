from pathlib import Path

from apply_pilot import resume

SAMPLE = Path(__file__).parents[1] / "examples" / "sample_resume.md"


def test_markdown_resume_to_profile():
    p = resume.ingest(SAMPLE)
    assert p["name"] == "Riley Quinn"
    assert p["email"] == "riley.quinn@example.com"
    assert p["phone"] == "(555) 010-0142"
    for s in ["Python", "TypeScript", "Go", "PostgreSQL", "Docker", "FastAPI", "AWS"]:
        assert s in p["skills"], s
    assert "Kubernetes" not in p["skills"]
    assert len(p["experience"]) == 2 and len(p["experience"][0]["bullets"]) == 3
    assert any("p95 query latency" in f for f in p["facts"])


def test_ambiguous_aliases_are_case_sensitive():
    assert "Go" not in resume.find_skills("Help us go to market faster")
    assert "Go" in resume.find_skills("Services written in Go and Python")
    assert "C++" in resume.find_skills("Strong C++ skills") and "C" not in resume.find_skills("C++")


def test_docx_and_pdf_ingestion(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Sam Example", 0)
    d.add_heading("Experience", 1)
    d.add_paragraph("Intern, Example Co")
    d.add_paragraph("Built Kafka pipelines in Java", style="List Bullet")
    d.save(tmp_path / "r.docx")
    p = resume.ingest(tmp_path / "r.docx")
    assert p["name"] == "Sam Example" and "Kafka" in p["skills"]
    assert p["experience"][0]["bullets"] == ["Built Kafka pipelines in Java"]

    # PDF: write a minimal one-page PDF by hand (no reportlab dependency)
    text = "Pat Example  Experience  Wrote Rust services on Kubernetes"
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    (tmp_path / "r.pdf").write_bytes(out)
    p = resume.ingest(tmp_path / "r.pdf")
    assert "Rust" in p["skills"] and "Kubernetes" in p["skills"]
